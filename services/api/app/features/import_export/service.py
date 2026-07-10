from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from typing import Any

import xlrd
from botocore.exceptions import ClientError
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.storage import (
    StorageNotConfiguredError,
    get_receipts_bucket,
    get_s3_client,
)
from app.features.attachments.models import ExpenseAttachment
from app.features.categories.models import UserCategory
from app.features.currency.service import (
    convert_to_base_or_unconverted,
    month_key_for,
    resolve_base_currency,
)
from app.features.exchange_rates.service import resolve_exchange_rate
from app.features.expenses.models import Expense
from app.features.expenses.schemas import (
    _MAX_AMOUNT,
    _validate_amount,
    _validate_currency,
)
from app.features.import_export.schemas import ImportSummary, PendingDeletion
from app.features.recurring.service import project_month

EXPORT_COLUMNS = [
    "Row ID",
    "Row Type",
    "Category",
    "Description",
    "Amount",
    "Currency",
    "Date",
    "Category ID",
    "Recurring Expense ID",
    "Installment Group ID",
    "Installment Index",
    "Installment Total",
    "Original Description",
    "Base Amount",
    "Exchange Rate",
]

EXPORT_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _year_bounds(year: int) -> tuple[date, date]:
    return date(year, 1, 1), date(year, 12, 31)


def _category_name_map(db: Session, user_id: str) -> dict[str, str]:
    rows = db.execute(
        select(UserCategory.id, UserCategory.name).where(
            UserCategory.user_id == user_id,
            UserCategory.household_id.is_(None),
        )
    ).all()
    return {category_id: name for category_id, name in rows}


def _recorded_rows(
    db: Session, user_id: str, year: int, categories: dict[str, str]
) -> list[dict[str, Any]]:
    year_start, year_end = _year_bounds(year)
    query = (
        select(Expense)
        .where(
            Expense.user_id == user_id,
            Expense.household_id.is_(None),
            Expense.spent_on.between(year_start, year_end),
        )
        .order_by(Expense.spent_on, Expense.id)
    )
    return [
        {
            "row_id": expense.id,
            "row_type": "recorded",
            "category": categories.get(expense.category_id),
            "description": expense.description,
            "amount": expense.amount,
            "currency": expense.currency,
            "date": expense.spent_on,
            "category_id": expense.category_id,
            "recurring_expense_id": expense.recurring_expense_id,
            "installment_group_id": expense.installment_group_id,
            "installment_index": expense.installment_index,
            "installment_total": expense.installment_total,
            "original_description": expense.original_description,
            "base_amount": expense.base_amount,
            "exchange_rate": expense.exchange_rate,
        }
        for expense in db.execute(query).scalars().all()
    ]


def _projected_rows(
    db: Session,
    user_id: str,
    year: int,
    today: date,
    categories: dict[str, str],
) -> list[dict[str, Any]]:
    """Recurring occurrences for months the monthly generation job hasn't reached yet.

    The worker only materializes the current month's rows (on the 1st, PRD
    §7.3), so a year's remaining future months have no Expense rows yet --
    project them so the export still reflects the full year.
    """
    current_month_start = today.replace(day=1)
    rows = []
    for month in range(1, 13):
        if date(year, month, 1) <= current_month_start:
            continue
        for occurrence in project_month(db, user_id, year, month):
            rows.append(
                {
                    "row_id": None,
                    "row_type": "projected",
                    "category": categories.get(occurrence.category_id),
                    "description": occurrence.description,
                    "amount": occurrence.amount,
                    "currency": occurrence.currency,
                    "date": occurrence.spent_on,
                    "category_id": occurrence.category_id,
                    "recurring_expense_id": occurrence.recurring_expense_id,
                    "installment_group_id": None,
                    "installment_index": None,
                    "installment_total": None,
                    "original_description": None,
                    # A projected row is a preview, never persisted, so there's
                    # no stored conversion to show -- left blank rather than
                    # computed, since import always ignores "projected" rows
                    # anyway (see build_import_plan).
                    "base_amount": None,
                    "exchange_rate": None,
                }
            )
    return rows


def build_year_export_rows(
    db: Session, user_id: str, year: int, today: date | None = None
) -> list[dict[str, Any]]:
    today = today or date.today()
    categories = _category_name_map(db, user_id)
    rows = _recorded_rows(db, user_id, year, categories)
    rows.extend(_projected_rows(db, user_id, year, today, categories))
    rows.sort(key=lambda row: (row["date"], row["row_id"] or ""))
    return rows


def build_export_workbook(
    db: Session, user_id: str, year: int, today: date | None = None
) -> BytesIO:
    rows = build_year_export_rows(db, user_id, year, today)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"
    sheet.append(EXPORT_COLUMNS)
    for row in rows:
        sheet.append(
            [
                row["row_id"],
                row["row_type"],
                row["category"],
                row["description"],
                str(row["amount"]),
                row["currency"],
                row["date"],
                row["category_id"],
                row["recurring_expense_id"],
                row["installment_group_id"],
                row["installment_index"],
                row["installment_total"],
                row["original_description"],
                str(row["base_amount"]) if row["base_amount"] is not None else None,
                str(row["exchange_rate"]) if row["exchange_rate"] is not None else None,
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def export_filename(year: int, today: date | None = None) -> str:
    today = today or date.today()
    return f"expenses-{year}-{today.isoformat()}.xlsx"


# --- Import ---

IMPORT_REQUIRED_COLUMNS = {
    "Row ID",
    "Row Type",
    "Category",
    "Description",
    "Amount",
    "Currency",
    "Date",
}

# (column header, Expense attribute) -- these are never written by import; a row
# that changes one of them relative to the current DB value is rejected so a
# spreadsheet edit can't silently corrupt installment/recurring provenance
# (PRD §13 risk).
SYSTEM_LOCKED_COLUMNS = [
    ("Recurring Expense ID", "recurring_expense_id"),
    ("Installment Group ID", "installment_group_id"),
    ("Installment Index", "installment_index"),
    ("Installment Total", "installment_total"),
    ("Original Description", "original_description"),
]


class WorkbookParseError(Exception):
    pass


class ImportValidationError(Exception):
    def __init__(self, errors: list[dict[str, Any]]) -> None:
        self.errors = errors
        super().__init__("Import validation failed")


class ImportRequiresConfirmationError(Exception):
    """Raised instead of committing when the sheet implies deletions the
    caller hasn't confirmed yet -- see `import_workbook`."""

    def __init__(self, deletions: list[PendingDeletion]) -> None:
        self.deletions = deletions
        super().__init__("Import requires confirmation of pending deletions")


@dataclass
class _InsertPlan:
    category_id: str
    description: str
    amount: Decimal
    currency: str
    spent_on: date
    # Sheet-supplied conversion, already validated against `amount` -- None
    # for either means the sheet left both blank, so apply_import_plan
    # computes them instead (see _convert_for_import).
    base_amount: Decimal | None = None
    exchange_rate: Decimal | None = None


@dataclass
class _UpdatePlan:
    expense: Expense
    category_id: str
    description: str
    amount: Decimal
    currency: str
    spent_on: date
    base_amount: Decimal | None = None
    exchange_rate: Decimal | None = None


def _clean_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _blank(values: list[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def _validate_header(header: list[str]) -> None:
    missing = IMPORT_REQUIRED_COLUMNS - set(header)
    if missing:
        raise WorkbookParseError(
            f"Missing required columns: {', '.join(sorted(missing))}"
        )
    seen = set()
    duplicates = set()
    for label in header:
        if label in seen:
            duplicates.add(label)
        seen.add(label)
    if duplicates:
        raise WorkbookParseError(f"Duplicate columns: {', '.join(sorted(duplicates))}")


def _rows_from_matrix(
    header: list[str], data_rows: list[tuple[int, list[Any]]]
) -> list[dict[str, Any]]:
    rows = []
    for row_number, values in data_rows:
        if _blank(values):
            continue
        record = dict(zip(header, values))
        record["_row_number"] = row_number
        rows.append(record)
    return rows


def _read_xlsx_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise WorkbookParseError(
            "Could not read the uploaded file as an Excel workbook."
        ) from exc

    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise WorkbookParseError("The workbook has no header row.")
    header = [_clean_str(cell) for cell in header_row]
    _validate_header(header)

    data_rows = [
        (row_number, list(values))
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        )
    ]
    return _rows_from_matrix(header, data_rows)


# xlrd cell types with no meaningful "value" for our purposes -- treated as blank
# so a stray formula error/boolean cell fails normal field validation instead of
# silently importing xlrd's raw internal code (e.g. an error code as an Amount).
_XLS_BLANK_CTYPES = {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR, xlrd.XL_CELL_BOOLEAN}


def _xls_cell_value(sheet: Any, book: Any, row_index: int, col: int) -> Any:
    cell = sheet.cell(row_index, col)
    if cell.ctype in _XLS_BLANK_CTYPES:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode))
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return int(cell.value) if cell.value == int(cell.value) else cell.value
    return cell.value


def _read_xls_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    try:
        book = xlrd.open_workbook(file_contents=file_bytes)
    except Exception as exc:
        raise WorkbookParseError(
            "Could not read the uploaded file as an Excel workbook."
        ) from exc

    sheet = book.sheet_by_index(0)
    if sheet.nrows == 0:
        raise WorkbookParseError("The workbook has no header row.")
    header = [_clean_str(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
    _validate_header(header)

    data_rows = [
        (
            row_index + 1,
            [
                _xls_cell_value(sheet, book, row_index, col)
                for col in range(sheet.ncols)
            ],
        )
        for row_index in range(1, sheet.nrows)
    ]
    return _rows_from_matrix(header, data_rows)


def read_import_rows(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "xlsx":
        return _read_xlsx_rows(file_bytes)
    if suffix == "xls":
        return _read_xls_rows(file_bytes)
    raise WorkbookParseError("Unsupported file type; only .xlsx and .xls are accepted.")


def _owned_categories_by_name(db: Session, user_id: str) -> dict[str, str]:
    rows = db.execute(
        select(UserCategory.name, UserCategory.id).where(
            UserCategory.user_id == user_id,
            UserCategory.household_id.is_(None),
        )
    ).all()
    return {name: category_id for name, category_id in rows}


def _owned_expenses_for_year(db: Session, user_id: str, year: int) -> list[Expense]:
    year_start, year_end = _year_bounds(year)
    return list(
        db.execute(
            select(Expense).where(
                Expense.user_id == user_id,
                Expense.household_id.is_(None),
                Expense.spent_on.between(year_start, year_end),
            )
        )
        .scalars()
        .all()
    )


def _pending_deletions(db: Session, delete_ids: set[str]) -> list[PendingDeletion]:
    if not delete_ids:
        return []
    rows = db.execute(
        select(Expense.id, Expense.description, Expense.spent_on).where(
            Expense.id.in_(delete_ids)
        )
    ).all()
    return [
        PendingDeletion(row_id=row_id, description=description, spent_on=spent_on)
        for row_id, description, spent_on in rows
    ]


_AMOUNT_QUANT = Decimal("0.01")


def _parse_decimal(value: Any) -> Decimal | None:
    """Raises InvalidOperation for a value that can't be read as a number."""
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int | float):
        return Decimal(str(value))
    return Decimal(str(value).strip())


def _parse_finite_decimal(value: Any) -> Decimal | None:
    """Like `_parse_decimal`, but also rejects NaN/Infinity.

    `Decimal("nan")` and `Decimal("inf")` parse without raising
    `InvalidOperation`, but NaN raises on comparison (e.g. `<= 0`) and
    Infinity raises on `quantize()` -- both would otherwise slip past a
    parse step's `except InvalidOperation` and surface as an unhandled 500
    instead of a validation error.
    """
    parsed = _parse_decimal(value)
    if not parsed.is_finite():
        raise InvalidOperation("value must be finite")
    return parsed


def _parse_amount(value: Any) -> tuple[Decimal | None, str | None]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "Amount is required."
    try:
        amount = _parse_finite_decimal(value)
    except InvalidOperation:
        return None, "Amount must be a valid number."
    try:
        return _validate_amount(amount), None
    except ValueError as exc:
        return None, str(exc)


def _parse_currency(value: Any) -> tuple[str | None, str | None]:
    text = _clean_str(value)
    if not text:
        return None, "Currency is required."
    try:
        return _validate_currency(text), None
    except ValueError as exc:
        return None, str(exc)


_EXCHANGE_RATE_QUANT = Decimal("0.000001")
# Expense.exchange_rate is Numeric(18, 6) -- 12 digits before the point.
_MAX_EXCHANGE_RATE = Decimal("999999999999.999999")


def _parse_base_amount(value: Any) -> tuple[Decimal | None, str | None]:
    """Optional -- a blank cell means "let the system compute it" (see
    apply_import_plan). A non-blank cell must be a valid positive amount
    that fits Expense.base_amount (Numeric(12, 2), same limit as Amount) --
    otherwise a value that passes this and the Amount x Exchange Rate
    consistency check would still blow up at DB commit as a numeric
    overflow instead of an ImportValidationError. Whether it's *consistent*
    with Amount x Exchange Rate is checked by the caller once all three
    fields are parsed."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    try:
        amount = _parse_finite_decimal(value)
    except InvalidOperation:
        return None, "Base Amount must be a valid number."
    if amount <= 0:
        return None, "Base Amount must be positive."
    try:
        quantized = amount.quantize(_AMOUNT_QUANT, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None, "Base Amount exceeds maximum allowed value."
    if quantized > _MAX_AMOUNT:
        return None, "Base Amount exceeds maximum allowed value."
    return quantized, None


def _parse_exchange_rate(value: Any) -> tuple[Decimal | None, str | None]:
    """Optional, same blank-means-compute convention as Base Amount. Also
    quantized/capped to fit Expense.exchange_rate (Numeric(18, 6)) for the
    same reason -- a validated plan must never fail at DB commit."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    try:
        rate = _parse_finite_decimal(value)
    except InvalidOperation:
        return None, "Exchange Rate must be a valid number."
    if rate <= 0:
        return None, "Exchange Rate must be positive."
    try:
        quantized = rate.quantize(_EXCHANGE_RATE_QUANT, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None, "Exchange Rate exceeds maximum allowed value."
    if quantized > _MAX_EXCHANGE_RATE:
        return None, "Exchange Rate exceeds maximum allowed value."
    return quantized, None


def _parse_spent_on(value: Any) -> tuple[date | None, str | None]:
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    text = _clean_str(value)
    if not text:
        return None, "Date is required."
    try:
        return date.fromisoformat(text), None
    except ValueError:
        return None, "Date must be a valid date (YYYY-MM-DD)."


def build_import_plan(
    db: Session, user_id: str, rows: list[dict[str, Any]], year: int
) -> tuple[list[_InsertPlan], list[_UpdatePlan], set[str], list[dict[str, Any]]]:
    """Diff uploaded rows against current DB state for `year`.

    The sheet represents the full desired state for that year: a blank Row ID
    is an insert, a present Row ID is an update, and any of the user's
    existing rows for `year` whose Row ID never appears in the upload is a
    deletion (this is how a whole installment group gets deleted too -- by
    omitting every one of its rows). "projected" rows are recurring-worker
    previews, never real data, so they're always ignored.

    Only rows whose *current* spent_on falls in `year` are eligible update/
    delete targets -- a Row ID for an expense that lives in a different year
    (e.g. pasted in from another year's export by mistake) is rejected as
    unknown rather than silently edited through this year's import. Likewise,
    every row's Date must fall in `year` -- an out-of-year date is rejected
    rather than silently inserting/moving an expense into a year outside the
    scope this import is diffing against.
    """
    categories = _owned_categories_by_name(db, user_id)
    by_id = {
        expense.id: expense for expense in _owned_expenses_for_year(db, user_id, year)
    }
    universe_ids = set(by_id)

    inserts: list[_InsertPlan] = []
    updates: list[_UpdatePlan] = []
    seen_ids: set[str] = set()
    errors: list[dict[str, Any]] = []

    for row in rows:
        row_number = row["_row_number"]
        if _clean_str(row.get("Row Type")).lower() == "projected":
            continue

        messages: list[str] = []

        description = _clean_str(row.get("Description"))
        if not description:
            messages.append("Description is required.")

        amount, amount_error = _parse_amount(row.get("Amount"))
        if amount_error:
            messages.append(amount_error)

        currency, currency_error = _parse_currency(row.get("Currency"))
        if currency_error:
            messages.append(currency_error)

        spent_on, date_error = _parse_spent_on(row.get("Date"))
        if date_error:
            messages.append(date_error)
        elif spent_on.year != year:
            messages.append(f"Date must be in {year}.")

        base_amount, base_amount_error = _parse_base_amount(row.get("Base Amount"))
        if base_amount_error:
            messages.append(base_amount_error)

        exchange_rate, exchange_rate_error = _parse_exchange_rate(
            row.get("Exchange Rate")
        )
        if exchange_rate_error:
            messages.append(exchange_rate_error)

        if not base_amount_error and not exchange_rate_error:
            if (base_amount is None) != (exchange_rate is None):
                messages.append(
                    "Base Amount and Exchange Rate must both be provided, or "
                    "both left blank."
                )
            elif amount is not None and base_amount is not None:
                expected_base_amount = (amount * exchange_rate).quantize(
                    _AMOUNT_QUANT, rounding=ROUND_HALF_UP
                )
                if base_amount != expected_base_amount:
                    messages.append(
                        "Base Amount does not match Amount x Exchange Rate."
                    )

        category_name = _clean_str(row.get("Category"))
        category_id = categories.get(category_name) if category_name else None
        if not category_name:
            messages.append("Category is required.")
        elif category_id is None:
            messages.append(f"Unknown category '{category_name}'.")

        row_id = _clean_str(row.get("Row ID")) or None

        if row_id is None:
            for label, _attr in SYSTEM_LOCKED_COLUMNS:
                if _clean_str(row.get(label)):
                    messages.append(f"{label} must be blank for a new row.")
            if messages:
                errors.append({"row": row_number, "messages": messages})
                continue
            inserts.append(
                _InsertPlan(
                    category_id,
                    description,
                    amount,
                    currency,
                    spent_on,
                    base_amount,
                    exchange_rate,
                )
            )
            continue

        existing = by_id.get(row_id)
        if existing is None:
            messages.append(f"Row ID '{row_id}' was not found.")
            errors.append({"row": row_number, "messages": messages})
            continue

        if row_id in seen_ids:
            messages.append(f"Row ID '{row_id}' is listed more than once.")

        for label, attr in SYSTEM_LOCKED_COLUMNS:
            if label not in row:
                continue
            sheet_value = _clean_str(row.get(label)) or None
            current_value = getattr(existing, attr)
            current_str = None if current_value is None else str(current_value)
            if sheet_value != current_str:
                messages.append(f"{label} cannot be changed (it is system-managed).")

        if messages:
            errors.append({"row": row_number, "messages": messages})
            continue

        seen_ids.add(row_id)
        updates.append(
            _UpdatePlan(
                existing,
                category_id,
                description,
                amount,
                currency,
                spent_on,
                base_amount,
                exchange_rate,
            )
        )

    delete_ids = universe_ids - seen_ids
    return inserts, updates, delete_ids, errors


def _convert_for_import(
    db: Session, user_id: str, amount: Decimal, currency: str, spent_on: date
) -> tuple[Decimal, Decimal]:
    """Import is personal-scope only and has no interactive moment to
    capture a rate per row (unlike the create/update endpoints). Falls back
    to a manually maintained rate for the pair (BUD-52), then to recording
    the row unconverted rather than failing the whole import when a rate
    would be required.

    Only called when a row's sheet doesn't supply a validated Base Amount +
    Exchange Rate pair (see `apply_import_plan`, which otherwise trusts
    those sheet values directly) -- i.e. a genuinely new row a user typed by
    hand, or a legacy sheet exported before these columns existed.
    """
    base_currency = resolve_base_currency(
        db, user_id, month_key_for(spent_on), household_id=None
    )
    exchange_rate = resolve_exchange_rate(
        db, user_id, None, currency, base_currency, None
    )
    return convert_to_base_or_unconverted(
        amount, currency, base_currency, exchange_rate=exchange_rate
    )


def apply_import_plan(
    db: Session,
    user_id: str,
    inserts: list[_InsertPlan],
    updates: list[_UpdatePlan],
    delete_ids: set[str],
) -> ImportSummary:
    object_keys: list[str] = []
    if delete_ids:
        # Captured before the delete -- the bulk DELETE below is raw SQL the
        # ORM can't hook, and expense_attachments cascades at the DB level,
        # so this is the only chance to know which S3 objects need cleanup.
        object_keys = list(
            db.execute(
                select(ExpenseAttachment.object_key).where(
                    ExpenseAttachment.expense_id.in_(delete_ids)
                )
            )
            .scalars()
            .all()
        )
        db.execute(
            delete(Expense).where(
                Expense.user_id == user_id,
                Expense.household_id.is_(None),
                Expense.id.in_(delete_ids),
            )
        )

    for update in updates:
        update.expense.category_id = update.category_id
        update.expense.description = update.description
        update.expense.amount = update.amount
        update.expense.currency = update.currency
        update.expense.spent_on = update.spent_on
        # A validated Base Amount + Exchange Rate on the sheet (already
        # checked consistent with Amount, in build_import_plan) is trusted
        # as-is -- this is what makes a plain export -> re-import round-trip
        # exactly, rather than re-deriving the conversion against whatever
        # base currency/manual rate happens to resolve today.
        if update.base_amount is not None and update.exchange_rate is not None:
            update.expense.base_amount = update.base_amount
            update.expense.exchange_rate = update.exchange_rate
        else:
            update.expense.base_amount, update.expense.exchange_rate = (
                _convert_for_import(
                    db, user_id, update.amount, update.currency, update.spent_on
                )
            )

    for insert in inserts:
        if insert.base_amount is not None and insert.exchange_rate is not None:
            base_amount, exchange_rate = insert.base_amount, insert.exchange_rate
        else:
            base_amount, exchange_rate = _convert_for_import(
                db, user_id, insert.amount, insert.currency, insert.spent_on
            )
        db.add(
            Expense(
                user_id=user_id,
                category_id=insert.category_id,
                description=insert.description,
                amount=insert.amount,
                currency=insert.currency,
                spent_on=insert.spent_on,
                base_amount=base_amount,
                exchange_rate=exchange_rate,
            )
        )

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise ImportValidationError(
            [
                {
                    "row": 0,
                    "messages": [
                        "Import could not be applied because it conflicts with "
                        "another row (e.g. a recurring occurrence already exists "
                        "for that date). Please review and re-upload."
                    ],
                }
            ]
        ) from exc

    # DB rows are already committed (source of truth for what's deleted), so
    # this best-effort cleanup mirrors attachments/service.py::delete_attachment:
    # a storage error here only orphans an object (a cost concern), not a
    # dangling row -- it must not surface as an error on an import that already
    # succeeded.
    if object_keys:
        s3 = get_s3_client()
        bucket = get_receipts_bucket()
        for object_key in object_keys:
            try:
                s3.delete_object(Bucket=bucket, Key=object_key)
            except (ClientError, StorageNotConfiguredError):
                pass

    return ImportSummary(
        inserted=len(inserts), updated=len(updates), deleted=len(delete_ids)
    )


def import_workbook(
    db: Session,
    user_id: str,
    file_bytes: bytes,
    filename: str,
    year: int,
    confirm_deletions: bool = False,
) -> ImportSummary:
    rows = read_import_rows(file_bytes, filename)
    inserts, updates, delete_ids, errors = build_import_plan(db, user_id, rows, year)
    if errors:
        raise ImportValidationError(errors)
    if delete_ids and not confirm_deletions:
        raise ImportRequiresConfirmationError(_pending_deletions(db, delete_ids))
    return apply_import_plan(db, user_id, inserts, updates, delete_ids)
