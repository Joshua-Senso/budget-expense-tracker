from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from app.features.expenses.models import Expense
from app.features.import_export.service import (
    EXPORT_COLUMNS,
    ImportValidationError,
    WorkbookParseError,
    _InsertPlan,
    _UpdatePlan,
    apply_import_plan,
    build_export_workbook,
    build_import_plan,
    build_year_export_rows,
    export_filename,
    import_workbook,
    read_import_rows,
)
from app.features.recurring.models import RecurringExpense


def _mock_db() -> MagicMock:
    return MagicMock()


def _make_expense(**kwargs) -> Expense:
    defaults = {
        "id": "exp-1",
        "user_id": "user-1",
        "category_id": "cat-1",
        "description": "Lunch",
        "amount": Decimal("150.00"),
        "currency": "PHP",
        "spent_on": date(2026, 1, 5),
        "installment_group_id": None,
        "installment_index": None,
        "installment_total": None,
        "original_description": None,
        "recurring_expense_id": None,
        "household_id": None,
    }
    exp = MagicMock(spec=Expense)
    for k, v in {**defaults, **kwargs}.items():
        setattr(exp, k, v)
    return exp


def _make_rule(**kwargs) -> RecurringExpense:
    defaults = {
        "id": "rec-1",
        "user_id": "user-1",
        "category_id": "cat-1",
        "description": "Netflix",
        "amount": Decimal("500.00"),
        "currency": "PHP",
        "start_on": date(2026, 1, 15),
        "frequency": "monthly",
        "is_active": True,
        "end_on": None,
        "household_id": None,
    }
    rule = MagicMock(spec=RecurringExpense)
    for k, v in {**defaults, **kwargs}.items():
        setattr(rule, k, v)
    return rule


@pytest.fixture(autouse=True)
def _default_base_currency(monkeypatch: pytest.MonkeyPatch) -> None:
    """Import rows default to "PHP" in these tests; default the resolved
    base currency to match so the same-currency short circuit applies for
    them. A row in a different currency has no interactive moment to supply
    a rate, so it falls back to unconverted (see `_convert_for_import`) --
    exercised explicitly by the currency-mismatch test below."""
    monkeypatch.setattr(
        "app.features.import_export.service.resolve_base_currency",
        lambda *args, **kwargs: "PHP",
    )


@pytest.fixture(autouse=True)
def _no_stored_exchange_rate_by_default(monkeypatch: pytest.MonkeyPatch) -> None:
    """Default the manual-rate lookup (BUD-52) to a pass-through -- no stored
    rate for any pair -- so existing tests don't need to know about it.
    Fallback-specific tests override this per-test."""
    monkeypatch.setattr(
        "app.features.import_export.service.resolve_exchange_rate",
        lambda db, currency, base_currency, exchange_rate: exchange_rate,
    )


def _queue_db(db: MagicMock, *result_sets) -> None:
    """Chain db.execute(...).all()/.scalars().all() return values in call order."""
    execute_results = []
    for result_set in result_sets:
        mock_result = MagicMock()
        mock_result.all.return_value = result_set
        mock_result.scalars.return_value.all.return_value = result_set
        execute_results.append(mock_result)
    db.execute.side_effect = execute_results


# --- build_year_export_rows ---


def test_build_year_export_rows_includes_recorded_expense() -> None:
    db = _mock_db()
    expense = _make_expense()
    _queue_db(
        db,
        [("cat-1", "Food")],  # category name map
        [expense],  # recorded expenses
    )

    rows = build_year_export_rows(db, "user-1", 2026, today=date(2026, 12, 31))

    assert len(rows) == 1
    row = rows[0]
    assert row["row_id"] == "exp-1"
    assert row["row_type"] == "recorded"
    assert row["category"] == "Food"
    assert row["amount"] == Decimal("150.00")
    assert row["date"] == date(2026, 1, 5)


def test_build_year_export_rows_projects_future_months_only() -> None:
    db = _mock_db()
    rule = _make_rule(start_on=date(2026, 1, 15))
    # category map, recorded (none), then one project_month() call per month
    # from Feb (today's month is Jan, so Feb..Dec are "future" -> 11 calls)
    future_month_results = [[rule]] + [[] for _ in range(10)]
    _queue_db(
        db,
        [("cat-1", "Subscriptions")],
        [],
        *future_month_results,
    )

    rows = build_year_export_rows(db, "user-1", 2026, today=date(2026, 1, 10))

    assert len(rows) == 1
    row = rows[0]
    assert row["row_id"] is None
    assert row["row_type"] == "projected"
    assert row["category"] == "Subscriptions"
    assert row["recurring_expense_id"] == "rec-1"
    assert row["date"] == date(2026, 2, 15)


def test_build_year_export_rows_does_not_project_current_or_past_months() -> None:
    db = _mock_db()
    _queue_db(db, [], [])

    rows = build_year_export_rows(db, "user-1", 2025, today=date(2026, 7, 1))

    assert rows == []
    # category map + recorded query only -- no project_month calls for a past year
    assert db.execute.call_count == 2


def test_build_year_export_rows_sorted_by_date() -> None:
    db = _mock_db()
    later = _make_expense(id="exp-later", spent_on=date(2026, 3, 1))
    earlier = _make_expense(id="exp-earlier", spent_on=date(2026, 1, 1))
    _queue_db(db, [("cat-1", "Food")], [later, earlier])

    rows = build_year_export_rows(db, "user-1", 2026, today=date(2026, 12, 31))

    assert [row["row_id"] for row in rows] == ["exp-earlier", "exp-later"]


# --- build_export_workbook ---


def test_build_export_workbook_writes_header_and_rows() -> None:
    db = _mock_db()
    expense = _make_expense()
    _queue_db(db, [("cat-1", "Food")], [expense])

    buffer = build_export_workbook(db, "user-1", 2026, today=date(2026, 12, 31))

    workbook = load_workbook(buffer)
    sheet = workbook["Expenses"]
    header = [cell.value for cell in sheet[1]]
    assert header == EXPORT_COLUMNS

    data_row = [cell.value for cell in sheet[2]]
    assert data_row[0] == "exp-1"
    assert data_row[1] == "recorded"
    assert data_row[2] == "Food"
    assert data_row[3] == "Lunch"
    assert data_row[4] == "150.00"
    assert data_row[5] == "PHP"


def test_build_export_workbook_preserves_exact_decimal_amount() -> None:
    """Amounts are written as text, not float, so re-import can't misread e.g. 9.99."""
    db = _mock_db()
    expense = _make_expense(amount=Decimal("9.99"))
    _queue_db(db, [("cat-1", "Food")], [expense])

    buffer = build_export_workbook(db, "user-1", 2026, today=date(2026, 12, 31))

    workbook = load_workbook(buffer)
    sheet = workbook["Expenses"]
    assert sheet[2][4].value == "9.99"


# --- export_filename ---


def test_export_filename_includes_year_and_current_date() -> None:
    assert (
        export_filename(2026, today=date(2026, 7, 5)) == "expenses-2026-2026-07-05.xlsx"
    )


# --- read_import_rows ---


def _build_xlsx(header: list[str], data_rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in data_rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_read_import_rows_parses_xlsx() -> None:
    file_bytes = _build_xlsx(
        EXPORT_COLUMNS,
        [
            ["exp-1", "recorded", "Food", "Lunch", "150.00", "PHP", date(2026, 1, 5)]
            + [None] * (len(EXPORT_COLUMNS) - 7)
        ],
    )

    rows = read_import_rows(file_bytes, "expenses.xlsx")

    assert len(rows) == 1
    assert rows[0]["Row ID"] == "exp-1"
    assert rows[0]["Description"] == "Lunch"
    assert rows[0]["_row_number"] == 2


def test_read_import_rows_skips_blank_rows() -> None:
    file_bytes = _build_xlsx(
        EXPORT_COLUMNS,
        [
            [None] * len(EXPORT_COLUMNS),
            ["exp-1", "recorded", "Food", "Lunch", "150.00", "PHP", date(2026, 1, 5)]
            + [None] * (len(EXPORT_COLUMNS) - 7),
        ],
    )

    rows = read_import_rows(file_bytes, "expenses.xlsx")

    assert len(rows) == 1
    assert rows[0]["_row_number"] == 3


def test_read_import_rows_rejects_missing_required_columns() -> None:
    file_bytes = _build_xlsx(["Row ID", "Description"], [["exp-1", "Lunch"]])

    with pytest.raises(WorkbookParseError, match="Missing required columns"):
        read_import_rows(file_bytes, "expenses.xlsx")


def test_read_import_rows_rejects_unreadable_file() -> None:
    with pytest.raises(WorkbookParseError, match="Could not read"):
        read_import_rows(b"not-a-real-workbook", "expenses.xlsx")


def test_read_import_rows_rejects_unsupported_extension() -> None:
    with pytest.raises(WorkbookParseError, match="Unsupported file type"):
        read_import_rows(b"whatever", "expenses.csv")


def test_read_import_rows_rejects_duplicate_columns() -> None:
    header = [*EXPORT_COLUMNS, "Amount"]
    file_bytes = _build_xlsx(header, [["exp-1", "recorded", "Food", "Lunch", "150.00"]])

    with pytest.raises(WorkbookParseError, match="Duplicate columns"):
        read_import_rows(file_bytes, "expenses.xlsx")


def _build_xls(header: list[str], data_rows: list[list]) -> bytes:
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Expenses")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for col, value in enumerate(header):
        sheet.write(0, col, value)
    for row_index, row in enumerate(data_rows, start=1):
        for col, value in enumerate(row):
            if isinstance(value, date):
                sheet.write(row_index, col, value, date_style)
            else:
                sheet.write(row_index, col, value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_read_import_rows_parses_legacy_xls() -> None:
    file_bytes = _build_xls(
        EXPORT_COLUMNS,
        [
            ["exp-1", "recorded", "Food", "Lunch", "150.00", "PHP", date(2026, 1, 5)]
            + [""] * (len(EXPORT_COLUMNS) - 7)
        ],
    )

    rows = read_import_rows(file_bytes, "expenses.xls")

    assert len(rows) == 1
    assert rows[0]["Row ID"] == "exp-1"
    assert rows[0]["Description"] == "Lunch"
    assert rows[0]["Date"] == datetime(2026, 1, 5)
    assert rows[0]["_row_number"] == 2


# --- build_import_plan ---


def _row(row_number: int = 2, **overrides) -> dict:
    defaults = {
        "Row ID": None,
        "Row Type": "recorded",
        "Category": "Food",
        "Description": "Lunch",
        "Amount": "150.00",
        "Currency": "PHP",
        "Date": date(2026, 1, 5),
        "Category ID": None,
        "Recurring Expense ID": None,
        "Installment Group ID": None,
        "Installment Index": None,
        "Installment Total": None,
        "Original Description": None,
        "_row_number": row_number,
    }
    defaults.update(overrides)
    return defaults


def test_build_import_plan_inserts_row_without_row_id() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    inserts, updates, delete_ids, errors = build_import_plan(
        db, "user-1", [_row()], 2026
    )

    assert errors == []
    assert updates == []
    assert delete_ids == set()
    assert len(inserts) == 1
    assert inserts[0].category_id == "cat-1"
    assert inserts[0].description == "Lunch"
    assert inserts[0].amount == Decimal("150.00")
    assert inserts[0].currency == "PHP"
    assert inserts[0].spent_on == date(2026, 1, 5)


def test_build_import_plan_updates_existing_row() -> None:
    db = _mock_db()
    expense = _make_expense(id="exp-1", category_id="cat-1", spent_on=date(2026, 1, 5))
    _queue_db(db, [("Food", "cat-1")], [expense])

    inserts, updates, delete_ids, errors = build_import_plan(
        db,
        "user-1",
        [_row(**{"Row ID": "exp-1", "Description": "Brunch"})],
        2026,
    )

    assert errors == []
    assert inserts == []
    assert delete_ids == set()
    assert len(updates) == 1
    assert updates[0].expense is expense
    assert updates[0].description == "Brunch"


def test_build_import_plan_deletes_rows_missing_from_upload() -> None:
    db = _mock_db()
    expense = _make_expense(id="exp-1", category_id="cat-1", spent_on=date(2026, 1, 5))
    _queue_db(db, [("Food", "cat-1")], [expense])

    inserts, updates, delete_ids, errors = build_import_plan(db, "user-1", [], 2026)

    assert errors == []
    assert inserts == []
    assert updates == []
    assert delete_ids == {"exp-1"}


def test_build_import_plan_deletes_whole_installment_group_when_all_rows_omitted() -> (
    None
):
    db = _mock_db()
    member_1 = _make_expense(
        id="exp-1",
        category_id="cat-1",
        spent_on=date(2026, 1, 5),
        installment_group_id="grp-1",
        installment_index=1,
        installment_total=2,
        original_description="TV",
    )
    member_2 = _make_expense(
        id="exp-2",
        category_id="cat-1",
        spent_on=date(2026, 2, 5),
        installment_group_id="grp-1",
        installment_index=2,
        installment_total=2,
        original_description="TV",
    )
    _queue_db(db, [("Food", "cat-1")], [member_1, member_2])

    inserts, updates, delete_ids, errors = build_import_plan(db, "user-1", [], 2026)

    assert errors == []
    assert delete_ids == {"exp-1", "exp-2"}


def test_build_import_plan_ignores_projected_rows() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    row = _row(**{"Row Type": "projected", "Row ID": None})
    inserts, updates, delete_ids, errors = build_import_plan(db, "user-1", [row], 2026)

    assert inserts == []
    assert updates == []
    assert delete_ids == set()
    assert errors == []


def test_build_import_plan_rejects_blank_description() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Description": "  "})], 2026
    )

    assert len(errors) == 1
    assert "Description is required." in errors[0]["messages"]


def test_build_import_plan_rejects_invalid_amount() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Amount": "not-a-number"})], 2026
    )

    assert len(errors) == 1
    assert "Amount must be a valid number." in errors[0]["messages"]


def test_build_import_plan_rejects_non_positive_amount() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(db, "user-1", [_row(**{"Amount": "-5"})], 2026)

    assert len(errors) == 1
    assert errors[0]["messages"] == ["amount must be positive"]


def test_build_import_plan_rejects_invalid_currency() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Currency": "philippine pesos"})], 2026
    )

    assert len(errors) == 1
    assert "currency must be one of" in errors[0]["messages"][0]


def test_build_import_plan_rejects_invalid_date() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Date": "not-a-date"})], 2026
    )

    assert len(errors) == 1
    assert "Date must be a valid date" in errors[0]["messages"][0]


def test_build_import_plan_rejects_date_outside_target_year() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Date": date(2027, 1, 1)})], 2026
    )

    assert len(errors) == 1
    assert errors[0]["messages"] == ["Date must be in 2026."]


def test_build_import_plan_rejects_unknown_category() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Category": "Mystery"})], 2026
    )

    assert len(errors) == 1
    assert "Unknown category 'Mystery'." in errors[0]["messages"]


def test_build_import_plan_rejects_unknown_row_id() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    _, _, _, errors = build_import_plan(
        db, "user-1", [_row(**{"Row ID": "ghost"})], 2026
    )

    assert len(errors) == 1
    assert "Row ID 'ghost' was not found." in errors[0]["messages"]


def test_build_import_plan_scopes_expense_lookup_to_the_target_year() -> None:
    """A Row ID belonging to a different year is rejected, not silently updated.

    build_import_plan only ever sees expenses the DB query already scoped to
    `year`, so a Row ID copied in from another year's export can't match here
    -- this asserts the query itself carries that year filter.
    """
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    build_import_plan(db, "user-1", [_row(**{"Row ID": "exp-other-year"})], 2026)

    expenses_query = db.execute.call_args_list[1][0][0]
    assert "spent_on BETWEEN" in str(expenses_query)


def test_build_import_plan_rejects_duplicate_row_id() -> None:
    db = _mock_db()
    expense = _make_expense(id="exp-1", category_id="cat-1", spent_on=date(2026, 1, 5))
    _queue_db(db, [("Food", "cat-1")], [expense])

    rows = [
        _row(row_number=2, **{"Row ID": "exp-1"}),
        _row(row_number=3, **{"Row ID": "exp-1"}),
    ]
    _, updates, _, errors = build_import_plan(db, "user-1", rows, 2026)

    assert len(updates) == 1
    assert len(errors) == 1
    assert errors[0]["row"] == 3
    assert "listed more than once" in errors[0]["messages"][0]


def test_build_import_plan_rejects_system_column_set_on_insert() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])

    row = _row(**{"Installment Group ID": "grp-new"})
    _, _, _, errors = build_import_plan(db, "user-1", [row], 2026)

    assert len(errors) == 1
    assert "Installment Group ID must be blank for a new row." in errors[0]["messages"]


def test_build_import_plan_rejects_system_column_changed_on_update() -> None:
    db = _mock_db()
    expense = _make_expense(
        id="exp-1",
        category_id="cat-1",
        spent_on=date(2026, 1, 5),
        installment_group_id="grp-1",
        installment_index=1,
        installment_total=2,
        original_description="TV",
    )
    _queue_db(db, [("Food", "cat-1")], [expense])

    row = _row(
        **{
            "Row ID": "exp-1",
            "Installment Group ID": "grp-1",
            "Installment Index": 99,
            "Installment Total": 2,
            "Original Description": "TV",
        }
    )
    _, _, _, errors = build_import_plan(db, "user-1", [row], 2026)

    assert len(errors) == 1
    assert errors[0]["messages"] == [
        "Installment Index cannot be changed (it is system-managed)."
    ]


def test_build_import_plan_allows_update_when_system_columns_untouched() -> None:
    db = _mock_db()
    expense = _make_expense(
        id="exp-1",
        category_id="cat-1",
        spent_on=date(2026, 1, 5),
        installment_group_id="grp-1",
        installment_index=1,
        installment_total=2,
        original_description="TV",
    )
    _queue_db(db, [("Food", "cat-1")], [expense])

    row = _row(
        **{
            "Row ID": "exp-1",
            "Installment Group ID": "grp-1",
            "Installment Index": 1,
            "Installment Total": 2,
            "Original Description": "TV",
        }
    )
    _, updates, _, errors = build_import_plan(db, "user-1", [row], 2026)

    assert errors == []
    assert len(updates) == 1


# --- apply_import_plan ---


def test_apply_import_plan_inserts_updates_and_deletes() -> None:
    db = _mock_db()
    expense = _make_expense(id="exp-1", category_id="cat-1", description="Old")
    _queue_db(db, [])  # the bulk-delete statement's execute() call

    inserts = [
        _InsertPlan("cat-1", "New expense", Decimal("10.00"), "PHP", date(2026, 1, 1))
    ]
    updates = [
        _UpdatePlan(
            expense, "cat-2", "Updated", Decimal("20.00"), "USD", date(2026, 2, 1)
        )
    ]
    delete_ids = {"exp-old"}

    summary = apply_import_plan(db, "user-1", inserts, updates, delete_ids)

    assert summary.inserted == 1
    assert summary.updated == 1
    assert summary.deleted == 1
    db.add.assert_called_once()
    added_expense = db.add.call_args[0][0]
    assert added_expense.description == "New expense"
    assert added_expense.base_amount == Decimal("10.00")
    assert added_expense.exchange_rate == Decimal("1")
    assert expense.category_id == "cat-2"
    assert expense.description == "Updated"
    assert expense.amount == Decimal("20.00")
    assert expense.currency == "USD"
    assert expense.spent_on == date(2026, 2, 1)
    # USD != the resolved base currency (PHP) and import has no interactive
    # moment to supply a rate, so it falls back to unconverted (BUD-54 will
    # address faithful cross-currency import conversion).
    assert expense.base_amount == Decimal("20.00")
    assert expense.exchange_rate == Decimal("1")
    db.commit.assert_called_once()


def test_apply_import_plan_uses_stored_manual_rate_for_cross_currency_row(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """BUD-52: import has no interactive moment to supply a rate (see the
    unconverted-fallback test above), but it can still use a manually
    maintained rate for the pair instead of falling all the way back to
    recording the row unconverted."""
    monkeypatch.setattr(
        "app.features.import_export.service.resolve_exchange_rate",
        lambda db, currency, base_currency, exchange_rate: Decimal("56.00"),
    )
    db = _mock_db()

    inserts = [
        _InsertPlan("cat-1", "New expense", Decimal("10.00"), "USD", date(2026, 1, 1))
    ]

    summary = apply_import_plan(db, "user-1", inserts, [], set())

    assert summary.inserted == 1
    added_expense = db.add.call_args[0][0]
    assert added_expense.base_amount == Decimal("560.00")
    assert added_expense.exchange_rate == Decimal("56.00")


def test_apply_import_plan_skips_delete_statement_when_nothing_to_delete() -> None:
    db = _mock_db()

    summary = apply_import_plan(db, "user-1", [], [], set())

    db.execute.assert_not_called()
    db.commit.assert_called_once()
    assert summary.inserted == 0
    assert summary.updated == 0
    assert summary.deleted == 0


def test_apply_import_plan_rolls_back_and_raises_on_db_conflict() -> None:
    from sqlalchemy.exc import IntegrityError

    db = _mock_db()
    db.commit.side_effect = IntegrityError("stmt", {}, Exception("conflict"))

    with pytest.raises(ImportValidationError) as exc_info:
        apply_import_plan(db, "user-1", [], [], set())

    db.rollback.assert_called_once()
    assert exc_info.value.errors[0]["row"] == 0


# --- import_workbook ---


def test_import_workbook_raises_validation_error_and_writes_nothing() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])
    file_bytes = _build_xlsx(
        EXPORT_COLUMNS,
        [
            [None, "recorded", "Food", "", "150.00", "PHP", date(2026, 1, 5)]
            + [None] * (len(EXPORT_COLUMNS) - 7)
        ],
    )

    with pytest.raises(ImportValidationError) as exc_info:
        import_workbook(db, "user-1", file_bytes, "expenses.xlsx", 2026)

    assert exc_info.value.errors[0]["row"] == 2
    db.add.assert_not_called()
    db.commit.assert_not_called()


def test_import_workbook_applies_valid_plan() -> None:
    db = _mock_db()
    _queue_db(db, [("Food", "cat-1")], [])
    file_bytes = _build_xlsx(
        EXPORT_COLUMNS,
        [
            [None, "recorded", "Food", "Lunch", "150.00", "PHP", date(2026, 1, 5)]
            + [None] * (len(EXPORT_COLUMNS) - 7)
        ],
    )

    summary = import_workbook(db, "user-1", file_bytes, "expenses.xlsx", 2026)

    assert summary.inserted == 1
    db.commit.assert_called_once()
